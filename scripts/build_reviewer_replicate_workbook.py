#!/usr/bin/env python3
"""Build the shareable reviewer workbook from the finished result tables.

Every number is read from a committed pipeline output; nothing is computed
here except the published-vs-new comparisons on the Validation sheets. Sources:

* replicate Tables 7/8/10: results/parallel_cross_source/replicate_reviewer_tables_all12
* replicate retrieval: results/replicate_retrieval_all12/tables
* raw cross-source Tables 4/5/6: results/parallel_cross_source/reviewer_raw_tables_all9
* dose sensitivity: results/parallel_cross_source/reviewer_final_summary
* published comparisons: results/parallel_cross_source/replicate_reviewer_tables_peer256_repaired,
  reviewer_final_summary/table5_*, and the raw Table 4/6 sheets of
  reviewer_manuscript_tables.xlsx (skipped if that workbook is absent)

Needs openpyxl. Usage::

    uv run --with openpyxl python scripts/build_reviewer_replicate_workbook.py
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
RESULTS = REPO / "results/parallel_cross_source"
T_DIR = RESULTS / "replicate_reviewer_tables_all12"
PUB_DIR = RESULTS / "replicate_reviewer_tables_peer256_repaired"
RAW_DIR = RESULTS / "reviewer_raw_tables_all9"
FINAL_DIR = RESULTS / "reviewer_final_summary"
R_DIR = REPO / "results/replicate_retrieval_all12/tables"
PUBLISHED_WORKBOOK = REPO / "reviewer_manuscript_tables.xlsx"

ORDER = ["l1000_phase1", "l1000_phase2", "tahoe", "cigs_mce", "novartis_batch_2500", "vcpi_0001",
         "cigs_tcm", "vcpi_0002", "gdpx2", "sciplex", "dilimap_train_val", "op3"]
LABEL = {"l1000_phase1": "L1000 Phase I", "l1000_phase2": "L1000 Phase II", "tahoe": "Tahoe-100M",
         "cigs_mce": "CIGS-MCE", "novartis_batch_2500": "Novartis/DRUG-seq U2OS", "vcpi_0001": "VCPI-0001",
         "cigs_tcm": "CIGS-TCM", "vcpi_0002": "VCPI-0002", "gdpx2": "GDPx2", "sciplex": "sci-Plex",
         "dilimap_train_val": "DILImap", "op3": "OP3"}
NEW = {"novartis_batch_2500", "vcpi_0001", "vcpi_0002", "gdpx2", "dilimap_train_val", "op3"}
PAIRS = [("cigs_mce", "cigs_tcm"), ("l1000_phase1", "cigs_mce"), ("l1000_phase1", "l1000_phase2"),
         ("l1000_phase1", "sciplex"), ("l1000_phase1", "tahoe"), ("l1000_phase2", "cigs_mce"),
         ("l1000_phase2", "sciplex"), ("l1000_phase2", "tahoe"), ("tahoe", "sciplex")]


def pair_label(a: str, b: str) -> str:
    return f"{LABEL[a]}–{LABEL[b]}"


def is_cigs_pair(a: str, b: str) -> bool:
    return a.startswith("cigs") or b.startswith("cigs")


FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="17365D")  # as in reviewer_manuscript_tables.xlsx
THIN = Side(style="thin", color="BFBFBF")


# ---------------------------------------------------------------- formatting
def est(mean, low, high, signed=False):
    if not all(np.isfinite(v) for v in (mean, low, high)):
        return ""
    fmt = "{:+.3f}" if signed else "{:.3f}"
    return f"{fmt.format(mean)} [{low:.3f}, {high:.3f}]"


def ci_cell(row, signed=False):
    if str(row["ci_status"]) == "zero_bootstrap_variance":
        return f"{row['mean']:.3f} [{row['mean']:.3f}, {row['mean']:.3f}]"
    return est(row["mean"], row["ci_low"], row["ci_high"], signed)


def parse_est(text):
    numbers = [float(x) for x in re.findall(r"[-+]?\d+\.\d+", str(text))]
    return numbers if len(numbers) == 3 else None


# ---------------------------------------------------------------- sheet helpers
def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        x = ws.cell(row=row, column=c)
        x.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        x.fill = HEADER_FILL
        x.alignment = Alignment(wrap_text=True, vertical="center")
        x.border = Border(bottom=THIN)


def write_table(ws, headers, rows, start_row=1, widths=None, freeze=True):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            x = ws.cell(row=i, column=j, value=v)
            x.font = Font(name=FONT, size=10)
            x.border = Border(bottom=THIN)
            if isinstance(v, float):
                x.number_format = "0.000"
            elif isinstance(v, int):
                x.number_format = "#,##0"
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=3 if len(headers) > 3 else 1)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[start_row].height = 30
    return start_row + len(rows)


def write_notes(ws, row, notes, ncol=8, title="Notes"):
    row += 2
    ws.cell(row=row, column=1, value=title).font = Font(name=FONT, bold=True, size=10)
    for n in notes:
        row += 1
        ws.cell(row=row, column=1, value=n).font = Font(name=FONT, size=10)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 28
    return row


# ---------------------------------------------------------------- replicate tables
STAT_HEADERS = ["Observed [95% CI]", "Centroid [95% CI]", "Observed minus centroid [95% CI]",
                "Individual-peer mean [95% CI]", "Observed minus peer [95% CI]", "Corrected percentile [95% CI]"]
STAT_SIGNED = [False, False, True, False, True, False]
REP_HEADERS = ["Representation", "Dataset", "Status", "n conditions", "n compounds", *STAT_HEADERS]
REP_WIDTHS = [30, 24, 11, 13, 12, 22, 22, 26, 24, 24, 24]


def metric_names(prefix, suffix, observed="observed"):
    return [f"{prefix}{stat}{suffix}" for stat in (
        observed, "centroid_baseline", "delta_vs_centroid", "individual_peer_mean",
        "delta_vs_individual_peer", "individual_peer_corrected_percentile")]


T_METRICS = ["tstat_observed_replicate_spearman", "tstat_centroid_baseline_spearman",
             "tstat_delta_vs_centroid_spearman", "tpeer_individual_peer_mean_spearman",
             "tpeer_delta_vs_individual_peer_spearman", "tpeer_individual_peer_corrected_percentile_spearman"]


def replicate_rows(ci, variants):
    """variants: (label, [six metric names in STAT_HEADERS order])."""
    rows = []
    for label, metrics in variants:
        for ds in ORDER:
            if (ds, metrics[0]) not in ci.index:
                continue
            obs = ci.loc[(ds, metrics[0])]
            row = [label, LABEL[ds], "New" if ds in NEW else "Rescored", int(obs["n_rows"]), int(obs["n_compounds"])]
            for metric, signed in zip(metrics, STAT_SIGNED):
                row.append(ci_cell(ci.loc[(ds, metric)], signed) if (ds, metric) in ci.index else "")
            rows.append(row)
    return rows


# ---------------------------------------------------------------- raw cross-source tables
def raw_metric_names(side, table):
    if table == 4:
        score, delta = "deg_lfc_spearman_pair_p05", "deg_lfc_spearman_p05"
        observed = "w4_observed_deg_lfc_spearman_sym_p05"
    else:
        score, delta = "spearman_logfc_pair", "spearman_logfc"
        observed = "w4_observed_spearman_logfc"
    return [observed, f"w4_{side}_centroid_{score}", f"w4_delta_vs_{side}_centroid_{delta}",
            f"w4_{side}_peer_{score}", f"w4_delta_vs_{side}_peer_{delta}",
            f"w4_{side}_peer_{score}_corrected_percentile"]


def raw_rows(ci, table):
    rows = []
    for a, b in PAIRS:
        for side in ("source", "target"):
            metrics = raw_metric_names(side, table)
            obs = ci.loc[(a, b, metrics[0])]
            row = [pair_label(a, b), side, "New" if is_cigs_pair(a, b) else "Rescored",
                   int(obs["n_compounds"]), int(obs["n_rows"])]
            for metric, signed in zip(metrics, STAT_SIGNED):
                row.append(ci_cell(ci.loc[(a, b, metric)], signed))
            rows.append(row)
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", type=Path, default=REPO / "reviewer_replicate_update_2026-09-24.xlsx")
    args = parser.parse_args()

    wb = Workbook()

    # ------------------------------------------------------------ Table 7
    t7 = pd.read_csv(T_DIR / "table_7_peer_baseline_cluster_bca_ci.tsv", sep="\t").set_index(["dataset_name", "metric"])
    rows7 = replicate_rows(t7, [
        ("raw logFC", metric_names("", "_deg_lfc_spearman_sym_p05")),
        ("z-score (dataset)", metric_names("dataset_normalized_", "_deg_lfc_spearman_sym_p05")),
        ("z-score (dataset x cell type)", metric_names("dataset_cell_type_normalized_", "_deg_lfc_spearman_sym_p05")),
    ])
    ws = wb.active
    ws.title = "Table 7 Replicates"
    end = write_table(ws, REP_HEADERS, rows7, widths=REP_WIDTHS)
    write_notes(ws, end, [
        "DEG-restricted replicate logFC Spearman (DEGs: adj.P.Value < 0.05). Under the z-score representations DEG membership stays on the raw adj.P.Value masks; only the ranked values are standardized, as in cross-dataset Table 4.",
        "z-score (dataset) = (logFC - mean) / SD per gene across the whole dataset; z-score (dataset x cell type) = the same within each dataset x cell line.",
        "Status: New = dataset added in this update; Rescored = previously reported dataset, rescored and checked against the published values (see Validation).",
    ], ncol=11)

    # ------------------------------------------------------------ Table 8
    t8 = pd.read_csv(T_DIR / "table_8_peer_baseline_cluster_bca_ci.tsv", sep="\t").set_index(["dataset_name", "metric"])
    rows8 = replicate_rows(t8, [("raw logFC sign", metric_names("", "_direction_agreement_p05"))])
    ws = wb.create_sheet("Table 8 Direction")
    end = write_table(ws, REP_HEADERS, rows8, widths=REP_WIDTHS)
    write_notes(ws, end, [
        "Replicate direction agreement on DEGs (adj.P.Value < 0.05). Kept on raw logFC by design: a per-gene z-score subtracts a population mean, which can flip a sign, so 'agreement in direction' would no longer mean agreement in biological direction.",
        "VCPI-0002 observed direction agreement is exactly 1.000 for every scored pair, so its interval has no width. It rests on 343 of 8,825 conditions: most VCPI-0002 conditions have too few DEGs for direction agreement to be defined.",
    ], ncol=11)

    # ------------------------------------------------------------ Table 10
    t10 = pd.read_csv(T_DIR / "table_10_peer_baseline_cluster_bca_ci.tsv", sep="\t").set_index(["dataset_name", "metric"])
    rows10 = replicate_rows(t10, [
        ("raw logFC, Spearman", metric_names("", "_spearman_logfc", "observed_replicate")),
        ("raw logFC, cosine", metric_names("raw_", "_cosine", "observed_replicate")),
        ("moderated t, Spearman", T_METRICS),
        ("z-score (dataset), Spearman", metric_names("dataset_normalized_", "_spearman_logfc", "observed_replicate")),
        ("z-score (dataset), cosine", metric_names("dataset_normalized_", "_cosine", "observed_replicate")),
        ("z-score (dataset x cell type), Spearman", metric_names("dataset_cell_type_normalized_", "_spearman_logfc", "observed_replicate")),
        ("z-score (dataset x cell type), cosine", metric_names("dataset_cell_type_normalized_", "_cosine", "observed_replicate")),
    ])
    ws = wb.create_sheet("Table 10 Replicates")
    end = write_table(ws, REP_HEADERS, rows10, widths=[40] + REP_WIDTHS[1:])
    write_notes(ws, end, [
        "All-gene replicate similarity: Spearman and cosine on raw logFC and both z-scores, and Spearman on the limma moderated t-statistic. New in this update: the z-score Spearman rows and the moderated-t individual-peer columns.",
        "Moderated t: the individual peers are the same sampled peers as for logFC, scored on the same genes. A condition with no same-dose other-compound peer has neither centroid nor peers, so those columns count slightly fewer conditions than n conditions (e.g. 2,437 of 2,517 for L1000 Phase I); the per-metric counts are on the Numeric sheet.",
        "L1000 observed replicate cosine (raw and both z-scores) now averages over every condition. The published value averaged only conditions with a centroid: raw cosine L1000 Phase I 0.1415 -> 0.1424, Phase II 0.1337 -> 0.1340; the z-score cosines move by at most 0.001. Restricted to the published conditions it reproduces 0.141544 exactly.",
    ], ncol=11)

    # ------------------------------------------------------------ Replicate retrieval
    rci = pd.read_csv(R_DIR / "replicate_retrieval_cluster_bca_ci.tsv", sep="\t").set_index(
        ["dataset_name", "similarity_metric", "scale_variant", "metric"])
    rtab = pd.read_csv(R_DIR / "replicate_retrieval_table.tsv", sep="\t")
    scale_label = {"raw": "raw logFC", "per_gene_population_zscore_dataset": "z-score (dataset)",
                   "per_gene_population_zscore_dataset_cell_type": "z-score (dataset x cell type)"}
    r_metrics = [("Observed rank [95% CI]", "observed_normalized_rank", False),
                 ("Recall@1 [95% CI]", "recall_at_1", False),
                 ("AUROC [95% CI]", "auroc", False),
                 ("Chance rank", "null_expected_normalized_rank", None),
                 ("Observed minus chance [95% CI]", "observed_minus_null", True),
                 ("Centroid rank [95% CI]", "centroid_normalized_rank", False),
                 ("Observed minus centroid [95% CI]", "delta_vs_centroid", True)]
    rrows = []
    for sim in ("cosine", "spearman"):
        for scale in scale_label:
            for ds in ORDER:
                t = rtab[(rtab.dataset_name == ds) & (rtab.similarity_metric == sim) & (rtab.scale_variant == scale)]
                if t.empty:
                    continue
                t = t.iloc[0]
                row = [sim.capitalize(), scale_label[scale], LABEL[ds], int(t.n_conditions), int(t.n_compounds),
                       int(round(t.mean_candidates))]
                for _, m, signed in r_metrics:
                    r = rci.loc[(ds, sim, scale, m)]
                    row.append(f"{r['mean']:.3f}" if signed is None else ci_cell(r, signed))
                rrows.append(row)
    ws = wb.create_sheet("Replicate Retrieval")
    end = write_table(ws, ["Similarity", "Representation", "Dataset", "n conditions", "n compounds", "Mean candidates",
                           *[h for h, _, _ in r_metrics]], rrows,
                      widths=[11, 26, 24, 13, 12, 14, 22, 22, 22, 12, 26, 22, 26])
    write_notes(ws, end, [
        "New analysis. Each replicate is a query; candidates are every other replicate in the same dataset x cell line x time (all doses and compounds); positives are the other replicates of the query's own condition. The query is never among its own candidates.",
        "Rank is the normalized best-positive rank, 1 - (best rank - 1) / (N - 1): 1 is best. Chance rank is its exact expectation with the positives placed at random; it exceeds 0.5 because queries have one or two positives. Recall@1 = a positive ranks first; AUROC = positives vs negatives.",
        "Centroid = mean of same-dose, other-compound replicates, injected into the candidate pool (rank among N + 1). No individual-peer baseline by design: within a dataset every other-compound replicate is already a candidate, so their mean rank is the middle of the pool by construction; the chance rank answers that question.",
        "The earlier replicate-retrieval code let a query retrieve its own sample (cyclic replicate pairing, no self-exclusion: 402 of 406 OP3 queries), which is why it reported ~0.997. It is retired; no earlier replicate-retrieval number should be reused.",
        "Strata need at least 10 compounds. 115 L1000 Phase II conditions sit in three 9-compound strata (CVCL_0031/6 h, CVCL_0332/3 h and 24 h) and are not scored.",
    ], ncol=13)

    # ------------------------------------------------------------ Raw Tables 4 and 6
    raw_ci = {}
    raw_headers = ["Dataset pair", "Peer side", "Status", "n compounds", "n matched rows", *STAT_HEADERS]
    raw_widths = [30, 10, 11, 12, 13, 22, 22, 26, 24, 24, 24]
    for table, title, what in ((4, "Table 4 Raw (9 pairs)", "DEG-restricted logFC Spearman (DEGs: adj.P.Value < 0.05)"),
                               (6, "Table 6 Raw (9 pairs)", "all-gene logFC Spearman")):
        ci = pd.read_csv(RAW_DIR / f"table{table}_raw_cluster_bca_ci.tsv", sep="\t")
        raw_ci[table] = ci
        ws = wb.create_sheet(title)
        end = write_table(ws, raw_headers, raw_rows(ci.set_index(["dataset_a", "dataset_b", "metric"]), table),
                          widths=raw_widths)
        write_notes(ws, end, [
            f"Cross-dataset {what} on raw logFC, for all nine dataset pairs. The three CIGS pairs are new; the published raw Table {table} had only the other six (raw CIGS rows were missing from the local cache then).",
            "Peer side: source = individual peers from the first dataset of the pair (same cell line, time and dose, different compound), scored against the second; target = the reverse. Centroid = the mean of those peers.",
            "Raw rows use up to 512 individual peers, as in the published raw Tables 4 and 6; the normalized tables use 256.",
            f"The six previously reported pairs reproduce the published point estimates exactly (see Validation Cross-dataset). The normalized Table {table} rows are unchanged from reviewer_manuscript_tables.xlsx.",
        ], ncol=11)

    # ------------------------------------------------------------ Dose sensitivity
    dose = pd.read_csv(FINAL_DIR / "dose_threshold_deg_metric_summary_complete.tsv", sep="\t")
    drows = []
    for a, b in (("cigs_mce", "cigs_tcm"), ("l1000_phase1", "cigs_mce"), ("l1000_phase2", "cigs_mce")):
        for _, r in dose[(dose.dataset_a == a) & (dose.dataset_b == b)].sort_values("threshold_order").iterrows():
            drows.append([pair_label(a, b), str(r.dose_threshold), int(r.n_matched_sample_pairs),
                          int(r.n_matching_conditions), int(r.n_matching_drugs), int(r.n_eligible_contexts),
                          float(r.mean_observed_deg_lfc_spearman_sym_p05), float(r.mean_baseline_pair_deg_lfc_spearman_p05),
                          float(r.mean_delta_vs_baseline_pair_deg_lfc_spearman_p05),
                          float(r.mean_observed_direction_agreement_p05), float(r.mean_baseline_pair_direction_agreement_p05),
                          float(r.mean_delta_vs_baseline_pair_direction_agreement_p05)])
    ws = wb.create_sheet("Dose Sensitivity CIGS")
    end = write_table(ws, ["Dataset pair", "Dose threshold", "n matched sample pairs", "n matching conditions",
                           "n matching compounds", "n eligible contexts", "DEG logFC Spearman: observed", "centroid",
                           "observed minus centroid", "Direction agreement: observed", "centroid", "observed minus centroid"],
                      drows, widths=[26, 14, 14, 14, 14, 12, 16, 12, 14, 16, 12, 14])
    write_notes(ws, end, [
        "The three CIGS pairs from dose_threshold_deg_metric_summary_complete.tsv (results/parallel_cross_source/reviewer_final_summary/), for comparison with your version. Thresholds: exact dose match, within 2x, within 3x, and the original 10x window (reference). Point estimates; the intervals are in dose_threshold_deg_cluster_bca_ci_complete.tsv.",
    ], ncol=12)

    # ------------------------------------------------------------ Validation (replicate)
    vrows = []
    for t in (7, 8, 10):
        old = pd.read_csv(PUB_DIR / f"table_{t}_peer_baseline_cluster_bca_ci.tsv", sep="\t")
        new = pd.read_csv(T_DIR / f"table_{t}_peer_baseline_cluster_bca_ci.tsv", sep="\t")
        j = old.merge(new, on=["dataset_name", "metric"], suffixes=("_pub", "_new"))
        for ds in [d for d in ORDER if d in set(j.dataset_name)]:
            x = j[j.dataset_name == ds]
            peer = x.metric.str.contains("individual_peer")
            vrows.append([f"Table {t}", LABEL[ds], int(x.n_rows_pub.iloc[0]), int(x.n_rows_new.iloc[0]),
                          "Yes" if (x.n_rows_pub == x.n_rows_new).all() else "No",
                          int((~peer).sum()), float(np.abs(x.mean_pub - x.mean_new)[~peer].max()),
                          int(peer.sum()), float(np.abs(x.mean_pub - x.mean_new)[peer].max())])
    ws = wb.create_sheet("Validation")
    end = write_table(ws, ["Table", "Dataset", "n conditions published", "n conditions now", "Counts match",
                           "Observed/centroid metrics compared", "Largest difference", "Peer metrics compared",
                           "Largest difference"], vrows, widths=[10, 18, 14, 14, 10, 16, 14, 14, 14])
    for r in range(2, end + 1):
        for c in (7, 9):
            ws.cell(row=r, column=c).number_format = "0.0E+00"
    write_notes(ws, end, [
        "Every previously reported dataset was rescored and compared with the published Tables 7/8/10 (replicate_reviewer_tables_peer256_repaired).",
        "CIGS-MCE and CIGS-TCM reproduce exactly, peers included: the published CIGS rows also used 256 peers with the same seed.",
        "Tahoe-100M also reproduces exactly, peers included. L1000 and sci-Plex: observed and centroid reproduce exactly; the peer columns differ by at most 8e-4 because their published rows used 512 peers and this update uses 256 throughout.",
        "The one other difference: L1000 Table 10 observed replicate cosine, whose published average counted only conditions with a centroid (up to 9e-4; see the Table 10 notes). Differences at the 1e-16 level are floating-point rounding.",
        "Moderated-t rows: no published values exist to compare with. The two independent scoring runs that both include sci-Plex and Tahoe give identical t-peer values for all 9,450 of their conditions, and every column the t-peer runs share with the main runs matches to within 1e-12.",
    ], ncol=9)

    # ------------------------------------------------------------ Validation (cross-dataset raw)
    xrows = []
    t5_pub = pd.read_csv(FINAL_DIR / "table5_direction_agreement_cluster_bca_ci.tsv", sep="\t")
    t5_new = pd.read_csv(RAW_DIR / "table5_direction_agreement_cluster_bca_ci.tsv", sep="\t")
    j5 = t5_pub.merge(t5_new, on=["dataset_a", "dataset_b", "metric"], suffixes=("_pub", "_new"))
    for a, b in PAIRS:
        x = j5[(j5.dataset_a == a) & (j5.dataset_b == b)]
        peer = x.metric.str.contains("peer")
        bounds = pd.concat([(x.ci_low_pub - x.ci_low_new).abs(), (x.ci_high_pub - x.ci_high_new).abs()])
        xrows.append(["Table 5", pair_label(a, b), "full precision",
                      "Yes" if ((x.n_rows_pub == x.n_rows_new) & (x.n_compounds_pub == x.n_compounds_new)).all() else "No",
                      int((~peer).sum()), float((x.mean_pub - x.mean_new).abs()[~peer].max()),
                      float(bounds[~pd.concat([peer, peer])].max()),
                      int(peer.sum()), float((x.mean_pub - x.mean_new).abs()[peer].max())])
    if PUBLISHED_WORKBOOK.is_file():
        published = load_workbook(PUBLISHED_WORKBOOK, read_only=True)
        reverse = {v: k for k, v in LABEL.items()}
        for table in (4, 6):
            ci = raw_ci[table].set_index(["dataset_a", "dataset_b", "metric"])
            by_pair: dict[tuple[str, str], list] = {}
            for r in published[f"Table {table} Raw"].iter_rows(min_row=2, values_only=True):
                if not r[0]:
                    continue
                a, b = (reverse[x] for x in r[0].split("–"))
                # published columns: n compounds, n rows, observed, peer mean, minus peer, corrected pct
                metrics = raw_metric_names(r[2], table)
                pairs = [(r[5], metrics[0]), (r[6], metrics[3]), (r[7], metrics[4]), (r[8], metrics[5])]
                entry = by_pair.setdefault((a, b), [0, 0, 0.0, True])
                obs = ci.loc[(a, b, metrics[0])]
                entry[3] &= (int(r[3]) == int(obs["n_compounds"])) and (int(r[4]) == int(obs["n_rows"]))
                for text, metric in pairs:
                    pub = parse_est(text)
                    new = ci.loc[(a, b, metric)]
                    entry[0] += 1
                    entry[1] += int(round(pub[0], 3) == round(float(new["mean"]), 3))
                    entry[2] = max(entry[2], abs(pub[1] - new["ci_low"]), abs(pub[2] - new["ci_high"]))
            for (a, b), (n, same, bound, counts) in by_pair.items():
                xrows.append([f"Table {table}", pair_label(a, b), "3 decimals (workbook)", "Yes" if counts else "No",
                              n, f"{same} of {n} identical", float(bound), "", ""])
    ws = wb.create_sheet("Validation Cross-dataset")
    end = write_table(ws, ["Table", "Dataset pair", "Published precision", "Counts match",
                           "Observed/centroid metrics compared", "Largest difference in estimate",
                           "Largest difference in a CI bound", "Peer metrics compared", "Largest peer difference"],
                      xrows, widths=[10, 30, 20, 10, 16, 18, 16, 14, 14])
    for r in range(2, end + 1):
        for c in (6, 7, 9):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, float):
                cell.number_format = "0.0E+00" if cell.value < 1e-3 else "0.0000"
    write_notes(ws, end, [
        "Table 5 (raw direction agreement) was published for all nine pairs, CIGS included, and rebuilt here from the same raw runs as the new Tables 4 and 6: observed and centroid estimates, their intervals and all counts match exactly, which validates the CIGS inputs. Its peer columns differ by at most 0.003 because the published Table 5 used 256 peers and these runs 512.",
        "Tables 4 and 6: the published raw sheets carry three decimals, so estimates are compared at that precision; for these rows 'Observed/centroid metrics compared' counts the four published columns per side (observed, peer mean, observed minus peer, corrected percentile). Every estimate and count is identical.",
        "The CI bounds of Tables 4 and 6 differ slightly (typically 0.001-0.005; 0.018 for Tahoe-100M–sci-Plex, 14 compounds). This is bootstrap Monte Carlo variation, not a data difference: the bootstrap seed is derived from a table label that differs from the original build, and changing only that label moves the bounds by the same amount.",
    ], ncol=9)

    # ------------------------------------------------------------ Numeric (for plotting)
    nrows = []
    for t in (7, 8, 10):
        ci = pd.read_csv(T_DIR / f"table_{t}_peer_baseline_cluster_bca_ci.tsv", sep="\t")
        for _, r in ci.iterrows():
            nrows.append([f"Table {t} replicate", LABEL[r.dataset_name], r.metric, float(r["mean"]), float(r.ci_low),
                          float(r.ci_high), int(r.n_rows), int(r.n_finite_rows), int(r.n_compounds), str(r.ci_status)])
    rc = pd.read_csv(R_DIR / "replicate_retrieval_cluster_bca_ci.tsv", sep="\t")
    for _, r in rc.iterrows():
        nrows.append(["Replicate retrieval", LABEL[r.dataset_name], f"{r.similarity_metric}|{r.scale_variant}|{r.metric}",
                      float(r["mean"]), float(r.ci_low), float(r.ci_high), int(r.n_rows), int(r.n_finite_rows),
                      int(r.n_compounds), str(r.ci_status)])
    for table, ci in raw_ci.items():
        for _, r in ci.iterrows():
            nrows.append([f"Table {table} raw", pair_label(r.dataset_a, r.dataset_b), r.metric, float(r["mean"]),
                          float(r.ci_low), float(r.ci_high), int(r.n_rows), int(r.n_finite_rows), int(r.n_compounds),
                          str(r.ci_status)])
    ws = wb.create_sheet("Numeric (for plotting)")
    end = write_table(ws, ["Table", "Dataset or pair", "Metric", "Mean", "CI low", "CI high", "n conditions or rows",
                           "n finite", "n compounds", "CI status"], nrows, widths=[20, 30, 70, 10, 10, 10, 13, 10, 12, 22])
    ws.auto_filter.ref = f"A1:J{end}"

    # ------------------------------------------------------------ Overview (first sheet)
    ws = wb.create_sheet("Overview", 0)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Chem-PerturBridge reviewer analysis: camera-ready update"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("24 September 2026. Replaces the replicate sheets and the raw Table 4/6 sheets of "
                "reviewer_manuscript_tables.xlsx; its other sheets (normalized Tables 4/6, Table 5, Table 9) are unchanged.")
    ws["A2"].font = Font(name=FONT, italic=True, size=10, color="595959")
    row = 4

    def section(title):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = Font(name=FONT, bold=True, size=11, color="17365D")
        row += 1

    def para(text, height=30):
        nonlocal row
        x = ws.cell(row=row, column=1, value=text)
        x.font = Font(name=FONT, size=10)
        x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.row_dimensions[row].height = height
        row += 1

    section("Your requests and where to find them")
    requests = [
        ("Peer baselines for replicate Tables 7, 8, 10, adding Novartis/DRUG-seq U2OS, VCPI-0001, VCPI-0002, GDPx2, DILImap, OP3",
         "Done (12 datasets)", "Tables 7, 8, 10 Replicates"),
        ("Normalized representations in Table 7, peers and centroid", "Done", "Table 7: both z-scores"),
        ("Normalized Spearman in Table 10, peers and centroid", "Done", "Table 10: z-score Spearman rows"),
        ("Which z-score normalization", "Both reported", "per Artur: dataset and dataset x cell type"),
        ("Minimal changes: 2 z-score representations on centroids, and peer baselines for Figure 2, for 12 datasets",
         "Done", "Tables 7, 10; Numeric (for plotting)"),
        ("Replicate cosine retrieval", "Done (new)", "Replicate Retrieval"),
        ("Peers on the moderated t-statistic (the deferred 'ideal variant')", "Done", "Table 10: moderated t rows"),
        ("L1000 Phase I/II condition counts differ from the published ones", "Resolved",
         "L1000 rescored with the published dataset set; counts reproduce exactly (2,517 / 2,729). See Validation."),
        ("Tables 4, 6 with the CIGS pairs: raw scale", "Done", "Table 4 Raw / Table 6 Raw (9 pairs)"),
        ("Tables 4, 6 with the CIGS pairs: normalized", "Done earlier", "reviewer_manuscript_tables.xlsx: Table 4/6 Normalized"),
        ("Dose sensitivity for the three CIGS pairs, to compare with yours", "Ours included", "Dose Sensitivity CIGS"),
        ("Where the scripts live", "See Code below", ""),
    ]
    for j, h in enumerate(["Request", "Status", "Where"], 1):
        ws.cell(row=row, column=j, value=h)
    style_header(ws, row, 3)
    row += 1
    for r in requests:
        for j, v in enumerate(r, 1):
            x = ws.cell(row=row, column=j, value=v)
            x.font = Font(name=FONT, size=10, bold=(j == 2))
            x.alignment = Alignment(wrap_text=True, vertical="top")
            x.border = Border(bottom=THIN)
        ws.row_dimensions[row].height = 30
        row += 1
    row += 1

    section("Settings")
    for text in [
        "Individual peers: same context and dose, different compound, deterministic sampling (seed 20260505); up to 256 per condition for the replicate tables, 512 for raw Tables 4 and 6 (as published). Centroids use every eligible peer.",
        "Intervals: compound-clustered BCa bootstrap, 2,000 iterations, resampling by PubChem CID.",
        "DEGs: adj.P.Value < 0.05. Replicate pairs are within one condition (same compound, cell line, time and dose).",
    ]:
        para(text)
    row += 1

    section("Notes for the methods / appendix")
    for text in [
        "Peer cap: every replicate row now uses 256 peers; the published L1000 and sci-Plex replicate rows used 512. This moves only their peer columns, by at most 8e-4.",
        "L1000 observed replicate cosine (Table 10) now averages over all conditions rather than only those with a centroid: raw cosine Phase I 0.1415 -> 0.1424.",
        "VCPI-0002 direction agreement (Table 8) is exactly 1.000, based on 343 of 8,825 conditions.",
        "Replicate retrieval is new. The earlier replicate-retrieval code let a query retrieve its own sample; no earlier replicate-retrieval number should be reused. 115 L1000 Phase II conditions in three 9-compound strata are below the 10-compound minimum.",
        "Replicates within a dataset can share plate or batch effects, which inflates any replicate-similarity measure (Tables 7, 8, 10 and retrieval alike).",
    ]:
        para(text)
    row += 1

    section("Code")
    for text in [
        "Branch cigs-production-hardening of github.com/elmella/Chem-PerturBridge_analysis; REVIEWER_ADDITIONS.md documents every step.",
        "scripts/run_replicate_jobs.sh reruns the replicate scoring, retrieval and t-peer jobs, resuming after an interruption; REVIEWER_ADDITIONS.md ('Camera-ready assembly') lists the commands that then build the tables and this workbook.",
        "scripts/replicate_reviewer_tables.py builds Tables 7, 8, 10; scripts/replicate_retrieval.py scores replicate retrieval; scripts/summarize_raw_cross_source_tables.py builds raw Tables 4, 5, 6; scripts/build_reviewer_replicate_workbook.py builds this workbook.",
    ]:
        para(text, height=30)

    wb.save(args.output)
    print(args.output)
    print("rows:", {"T7": len(rows7), "T8": len(rows8), "T10": len(rows10), "retrieval": len(rrows),
                    "dose": len(drows), "validation": len(vrows), "validation_cross": len(xrows), "numeric": len(nrows)})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
